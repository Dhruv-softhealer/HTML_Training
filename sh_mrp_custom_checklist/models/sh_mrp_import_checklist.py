# -*- coding: utf-8 -*-
# Part of Softhealer Technologies.

import base64
import csv
import io

import openpyxl
from odoo import Command, models, fields, api
from odoo.exceptions import UserError


class ImportChecklist(models.TransientModel):
    _name = 'sh.mrp.import.checklist'
    _description = 'MRP Import Checklist'
        
    
    import_type = fields.Selection([('csv', 'CSV File'), ('excel', 'EXCEL File')], string="Import File Type")
    company_id = fields.Many2one('res.company', string="Company", default=lambda self:self.env.company)
    file = fields.Binary(string="File")
    
    def apply_button(self):
        if not self.file:
            raise UserError(("Please upload a file to import."))
 
        if self.import_type == 'csv':
            self.import_csv()
        elif self.import_type == 'excel':
            self.import_excel()
 
        self.env["bus.bus"]._sendone(
                self.env.user.partner_id,
                "simple_notification",
                {
                    "type": "success",
                    "title": "Success",
                    "message": "Data successfully imported",},)
 
        
 
    def import_csv(self):
        try:
            data = base64.b64decode(self.file)
            file_input = io.StringIO(data.decode('utf-8'))
            reader = csv.DictReader(file_input)
        except Exception as e:
            raise UserError(("Invalid CSV file. Error: %s") % e)
 
        for row in reader:
            name = row.get('Name')
            description = row.get('Description')
            if not name:
                continue
            self.env['sh.mrp.custom.checklist'].create({
                'name': name,
                'description': description or '',
            })
 
    def import_excel(self):
        try:
            data = base64.b64decode(self.file)
            file_input = io.BytesIO(data)
            workbook = openpyxl.load_workbook(file_input)
            sheet = workbook.active
        except Exception as e:
            raise UserError(("Invalid Excel file. Error: %s") % e)
 
        headers = [cell.value for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
        for row in sheet.iter_rows(min_row=2, values_only=True):
            row_data = dict(zip(headers, row))
            name = row_data.get('name')
            description = row_data.get('description')
            if not name:
                continue
 
            self.env['sh.mrp.custom.checklist'].create({
                'name': name,
                'description': description or '',
            })
    
    # @api.onchange('file', 'import_type')
    # def check_file_type(self):
    #     if self.file and self.import_type:
    #         file_header = base64.b64decode(self.file)[:4]
 
    #         if self.import_type == 'csv' and file_header != b'name':  
    #             raise UserError(("You selected CSV, but the uploaded file is not a CSV file."))
    #         elif self.import_type == 'excel':
    #             if not (file_header.startswith(b'PK') or b'\xd0\xcf' in file_header):
    #                 raise UserError(("You selected Excel, but the uploaded file is not an Excel file."))
 
 